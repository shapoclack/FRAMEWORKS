"""Скрипт создания трёх отдельных официальных отчётов в формате DOCX по ПР5, ПР6 и ПР7."""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = Path(__file__).resolve().parent

IMG1_PATH = Path(r"C:/Users/frolo/.gemini/antigravity/brain/d57d169d-cf3c-4c46-b631-cd3105e21cb4/.user_uploaded/media_1792668603168_644a1c2f.png")
IMG2_PATH = Path(r"C:/Users/frolo/.gemini/antigravity/brain/d57d169d-cf3c-4c46-b631-cd3105e21cb4/.user_uploaded/media_1792668612473_35f68279.png")

STUDENT_NAME = "Фролов Кирилл Иванович"
STUDENT_GROUP = "ЭФБО 12-24"
PROJECT_THEME = "Система учета тем курсовых работ"
TEACHER = "Ящун Татьяна Викторовна"
SEMESTER = "5 семестр, 2026–2027 уч. год"
DISCIPLINE = "Технологии разработки приложений на базе фреймворков"


def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def create_base_doc(pr_number: int, pr_title: str) -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.79)
        section.right_margin = Inches(0.79)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Шапка института
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run(
        "МИНОБРНАУКИ РОССИИ\n"
        "РТУ МИРЭА — Российский технологический университет\n"
        "Институт перспективных технологий и индустриального программирования\n"
        "Кафедра индустриального программирования\n\n"
    )
    r_inst.font.size = Pt(9.5)
    r_inst.font.bold = True

    # Название работы
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_main = p_title.add_run(f"ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № {pr_number}\n")
    r_main.font.size = Pt(15)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    r_sub = p_title.add_run(f"«{pr_title}»\n")
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True

    r_disc = p_title.add_run(f"по дисциплине: «{DISCIPLINE}»\n\n")
    r_disc.font.size = Pt(11)
    r_disc.font.italic = True

    # Таблица общих сведений
    t_info = doc.add_table(rows=5, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_rows = [
        ("Студент:", STUDENT_NAME),
        ("Учебная группа:", STUDENT_GROUP),
        ("Тема индивидуального проекта:", PROJECT_THEME),
        ("Преподаватель:", TEACHER),
        ("Семестр:", SEMESTER),
    ]
    for idx, (label, val) in enumerate(info_rows):
        row = t_info.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.5)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FAFAFA")
        set_cell_margins(c0)
        set_cell_margins(c1)
        r0 = c0.paragraphs[0].add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r1 = c1.paragraphs[0].add_run(val)
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return doc


def add_section_heading(doc, text):
    h = doc.add_heading(level=2)
    r = h.add_run(text)
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(4)


# ==============================================================================
# ОТЧЕТ ПР5
# ==============================================================================
def generate_pr5_report():
    doc = create_base_doc(5, "DJANGO, HTTP И ОСНОВЫ ВЕБ-РАЗРАБОТКИ")

    add_section_heading(doc, "1. Цель работы")
    p = doc.add_paragraph(
        "Изучить назначение фреймворка Django, структуру Django-проекта и приложений, "
        "принципы обработки HTTP-запросов и формирования HTTP-ответов, маршрутизацию URL и основы "
        "создания веб-страниц. На основе объектной модели и JSON-хранилища индивидуального проекта, "
        "разработанных в ПР1–ПР3, подготовить проект к переходу от консольного приложения к веб-приложению на Django."
    )

    add_section_heading(doc, "2. Выполненные материалы Яндекс Практикума")
    doc.add_paragraph(
        "В рамках ПР5 полностью освоены материалы Раздела 04 «Основы веб-разработки» "
        "курса «Бэкенд-разработка» Яндекс Практикума:"
    )

    t = doc.add_table(rows=4, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["№", "Тема курса (Раздел 04)", "Статус"]):
        cell = t.rows[0].cells[i]
        set_cell_background(cell, "0D3B66")
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    yp_pr5 = [
        ("1", "Создание проекта («Введение в спринт», «Фреймворк Django», «Создание Django-проекта», «Приложения в Django-проекте»)", "Выполнено (100%)"),
        ("2", "Пути и view-функции («Протокол HTTP: стандарт для запросов и ответов», «Планирование адресов», «Конвертеры путей», «Распределение адресов»)", "Выполнено (100%)"),
        ("3", "Вёрстка для бэкендера («Веб-страницы: язык HTML», «Структура документа», «CSS стили и селекторы», «Bootstrap: быстрая верстка»)", "Выполнено (100%)"),
    ]
    for r_idx, vals in enumerate(yp_pr5, 1):
        for c_idx, val in enumerate(vals):
            cell = t.rows[r_idx].cells[c_idx]
            set_cell_background(cell, "F9FBFD" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(val)
            r.font.size = Pt(9.5)
            if c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x7A, 0x43)

    if IMG1_PATH.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.add_run().add_picture(str(IMG1_PATH), width=Inches(4.5))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = p_cap.add_run("Рисунок 1 — Подтверждение выполнения тем раздела 04 (Яндекс Практикум)")
        rc.font.italic = True
        rc.font.size = Pt(8.5)

    add_section_heading(doc, "3. Доработка индивидуального проекта")
    doc.add_paragraph(f"Название проекта: «{PROJECT_THEME}».")
    doc.add_paragraph("В ходе выполнения ПР5 в проекте реализовано:")
    points = [
        "Создан Django-проект coursework_project и два приложения: core (главная страница, информация) и coursework (предметная область тем курсовых работ);",
        "Настроена двухуровневая маршрутизация URL через include() и конвертеры путей <int:pk>;",
        "Реализованы первые view-функции (home, about, http_info, topic_list, topic_detail, student_list, supervisor_list);",
        "Настроен глобальный обработчик ошибок handler404 для возврата брендированной страницы «404 — Страница не найдена»;",
        "Подключен фреймворк Bootstrap 5.3 через CDN, создана адаптивная сетка и элементы интерфейса;",
        "Разработана демонстрационная страница /http-info/ с интерактивным разбором метода, пути, заголовков и query-параметров входящего HTTP-запроса;",
        "Сохранена связь с кодом ПР1–ПР3: консольная версия main.py, функции storage.py и JSON-файлы каталога data/ продолжают функционировать."
    ]
    for pt in points:
        doc.add_paragraph(pt, style='List Bullet')

    add_section_heading(doc, "4. Результат работы")
    doc.add_paragraph(
        "В результате выполнения ПР5 создано функционирующее веб-приложение Django с веб-интерфейсом на Bootstrap 5.3. "
        "Обеспечен сквозной цикл: клиентский HTTP-запрос → URLconf → View-функция → Подготовка данных → HTTP-ответ (HTML)."
    )

    add_section_heading(doc, "5. Контрольные проверки")
    doc.add_paragraph("• python manage.py check — 0 ошибок (System check identified no issues);\n"
                      "• pytest -v — 16 passed (тесты консольной версии ПР2–ПР3);\n"
                      "• flake8 . --extend-exclude=venv — 0 ошибок (соответствие PEP 8).")

    add_section_heading(doc, "6. Фиксация в Git")
    doc.add_paragraph("Результат зафиксирован коммитом с описанием перехода на Django и тегом: git tag PR5.")

    add_section_heading(doc, "7. Вывод")
    doc.add_paragraph(
        "В ходе практической работы № 5 успешно освоены основы веб-разработки на базе фреймворка Django, "
        "механизмы маршрутизации запросов и генерации HTTP-ответов. Проект переведён от консольного прототипа "
        "к базовому веб-интерфейсу с сохранением наработок предыдущих этапов."
    )

    out_file = BASE_DIR / "Отчет_ПР5_Фролов_Кирилл_Иванович.docx"
    doc.save(str(out_file))
    print(f"Отчет ПР5 создан: {out_file}")


# ==============================================================================
# ОТЧЕТ ПР6
# ==============================================================================
def generate_pr6_report():
    doc = create_base_doc(6, "ШАБЛОНЫ DJANGO И ПРОЕКТНАЯ РАЗРАБОТКА")

    add_section_heading(doc, "1. Цель работы")
    p = doc.add_paragraph(
        "Изучить систему шаблонов Django: базовый шаблон и наследование ({% extends %}, {% block %}), "
        "включаемые шаблоны ({% include %}), передачу данных из view-функций через контекст, циклы ({% for %}) "
        "и условные конструкции ({% if %}), фильтры шаблонов, именованные URL-маршруты с пространствами имён "
        "и подключение статических файлов (CSS, JS, изображения). Перевести веб-интерфейс индивидуального проекта "
        "с формирования HTML в Python на систему шаблонов Django (DTL)."
    )

    add_section_heading(doc, "2. Выполненные материалы Яндекс Практикума")
    doc.add_paragraph(
        "В рамках ПР6 освоены материалы уроков «HTML и шаблоны Django» и финальный проект спринта "
        "Раздела 04 «Основы веб-разработки» курса Яндекс Практикума:"
    )

    t = doc.add_table(rows=3, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["№", "Тема курса (Раздел 04)", "Статус"]):
        cell = t.rows[0].cells[i]
        set_cell_background(cell, "0D3B66")
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    yp_pr6 = [
        ("1", "HTML и шаблоны Django (теги include/extends/block, словарь контекста, ветвление, циклы, фильтры, подключение стилей, картинок и скриптов)", "Выполнено (100%)"),
        ("2", "Финальный проект спринта Блогикум (построение модульного веб-приложения на шаблонах DTL)", "Выполнено (100%)"),
    ]
    for r_idx, vals in enumerate(yp_pr6, 1):
        for c_idx, val in enumerate(vals):
            cell = t.rows[r_idx].cells[c_idx]
            set_cell_background(cell, "F9FBFD" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(val)
            r.font.size = Pt(9.5)
            if c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x7A, 0x43)

    if IMG1_PATH.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.add_run().add_picture(str(IMG1_PATH), width=Inches(4.5))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = p_cap.add_run("Рисунок 1 — Подтверждение выполнения тем по шаблонам Django (Яндекс Практикум)")
        rc.font.italic = True
        rc.font.size = Pt(8.5)

    add_section_heading(doc, "3. Доработка индивидуального проекта")
    doc.add_paragraph(f"Название проекта: «{PROJECT_THEME}».")
    doc.add_paragraph("В ходе выполнения ПР6 реализована шаблонная архитектура:")
    points = [
        "Создана структура каталогов templates/ уровня проекта (зарегистрирована в TEMPLATES['DIRS'] в settings.py) и внутри приложений core/templates/ и coursework/templates/;",
        "Создан единый базовый каркас templates/base.html с блоками {% block title %} и {% block content %}, адаптивным меню и подвалом сайта;",
        "Все view-функции переведены на вызов render(request, template, context) с полным удалением вспомогательной функции page();",
        "Созданы специализированные шаблоны: topic_list.html, topic_detail.html, student_list.html, supervisor_list.html, home.html, about.html, http_info.html;",
        "Реализованы конструкции DTL: циклы {% for item in items %} с веткой {% empty %}, условия {% if %}, фильтры данных (|length, |default, |date);",
        "Настроена именованная маршрутизация с пространствами имён (app_name = 'coursework', {% url 'coursework:topic_detail' topic.id %});",
        "Подключены статические файлы через {% load static %}: собственный файл static/css/style.css, иконки Bootstrap Icons."
    ]
    for pt in points:
        doc.add_paragraph(pt, style='List Bullet')

    add_section_heading(doc, "4. Результат работы")
    doc.add_paragraph(
        "Веб-часть проекта полностью переведена с формирования строк HTML в Python на систему шаблонов Django. "
        "Достигнуто разделение ответственности: Python-код готовит данные в словаре context, а HTML-документ "
        "формируется механизмом Django Template Language с поддержкой наследования и статических файлов."
    )

    add_section_heading(doc, "5. Контрольные проверки")
    doc.add_paragraph("• python manage.py check — 0 ошибок (System check identified no issues);\n"
                      "• pytest -v — 16 passed (тесты успешно проходят);\n"
                      "• flake8 . --extend-exclude=venv — 0 ошибок (соответствие PEP 8).")

    add_section_heading(doc, "6. Фиксация в Git")
    doc.add_paragraph("Изменения зафиксированы коммитом: git commit -m 'PR6: add Django templates and static files'.")

    add_section_heading(doc, "7. Вывод")
    doc.add_paragraph(
        "В ходе практической работы № 6 освоена система шаблонов Django. "
        "Реализовано разделение верстки и бизнес-логики, внедрен базовый шаблон base.html с наследованием блоков, "
        "настроена передача контекста и подключены собственные статические стили CSS."
    )

    out_file = BASE_DIR / "Отчет_ПР6_Фролов_Кирилл_Иванович.docx"
    doc.save(str(out_file))
    print(f"Отчет ПР6 создан: {out_file}")


# ==============================================================================
# ОТЧЕТ ПР7
# ==============================================================================
def generate_pr7_report():
    doc = create_base_doc(7, "ОСНОВЫ БАЗ ДАННЫХ И SQL")

    add_section_heading(doc, "1. Цель работы")
    p = doc.add_paragraph(
        "Изучить основные понятия реляционных баз данных и языка SQL, освоить работу с СУБД SQLite: "
        "проектирование структуры таблиц, выбор типов данных, определение первичных и внешних ключей, "
        "создание таблиц и выполнение запросов на выборку, фильтрацию, сортировку, ограничение данных, "
        "применение агрегирующих функций и группировки. Освоить получение данных из базы SQLite средствами Python "
        "и стандартной библиотеки sqlite3 (без Django ORM на данном этапе), подготовив переход индивидуального "
        "проекта от JSON-хранилища к реляционному хранению данных."
    )

    add_section_heading(doc, "2. Выполненные материалы Яндекс Практикума")
    doc.add_paragraph(
        "В рамках ПР7 выполнены материалы раздела 05 «Базы данных и Django ORM» курса Яндекс Практикума "
        "в части темы «Введение в базы данных»:"
    )

    t = doc.add_table(rows=11, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["№", "Тема / Урок (Раздел 05)", "Статус"]):
        cell = t.rows[0].cells[i]
        set_cell_background(cell, "0D3B66")
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    yp_pr7 = [
        ("1", "Введение в базы данных / Введение в спринт", "Выполнено"),
        ("2", "Базы данных: реляционные и не только", "Выполнено"),
        ("3", "SQL: язык структурированных запросов", "Выполнено"),
        ("4", "Первые запросы, первые результаты", "Выполнено"),
        ("5", "Работа с базой данных из Python (модуль sqlite3)", "Выполнено"),
        ("6", "Получение данных (SELECT, cursor, fetchall)", "Выполнено"),
        ("7", "Практика запросов", "Выполнено"),
        ("8", "Сортировка, ограничение и сдвиг выборки (ORDER BY, LIMIT, OFFSET)", "Выполнено"),
        ("9", "Агрегирующие функции (COUNT, SUM, AVG, MIN, MAX)", "Выполнено"),
        ("10", "Группировка записей в выборке (GROUP BY, фильтрация HAVING)", "Выполнено"),
    ]
    for r_idx, vals in enumerate(yp_pr7, 1):
        for c_idx, val in enumerate(vals):
            cell = t.rows[r_idx].cells[c_idx]
            set_cell_background(cell, "F9FBFD" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p_c = cell.paragraphs[0]
            r = p_c.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x7A, 0x43)

    if IMG2_PATH.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.add_run().add_picture(str(IMG2_PATH), width=Inches(4.5))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = p_cap.add_run("Рисунок 1 — Подтверждение выполнения темы «Введение в базы данных» (Яндекс Практикум)")
        rc.font.italic = True
        rc.font.size = Pt(8.5)

    add_section_heading(doc, "3. Доработка индивидуального проекта")
    doc.add_paragraph(f"Название проекта: «{PROJECT_THEME}».")

    doc.add_paragraph().add_run("3.1. Структура базы данных и схема (sql/schema.sql):").bold = True
    p_sch = doc.add_paragraph(
        "Создан файл базы данных data/coursework.sqlite. В файле sql/schema.sql созданы 4 таблицы:\n"
        "• students (id PK, full_name, group_name, direction, email UNIQUE, created_at DEFAULT datetime('now'));\n"
        "• supervisors (id PK, full_name, degree, position, email UNIQUE, max_topics DEFAULT 5 CHECK (1..20), created_at);\n"
        "• topics (id PK, title, description, deadline, created_at, status CHECK, supervisor_id FK REFERENCES supervisors(id) ON DELETE RESTRICT, student_id FK REFERENCES students(id) ON DELETE SET NULL);\n"
        "• status_history (id PK AUTOINCREMENT, topic_id FK REFERENCES topics(id) ON DELETE CASCADE, old_status, new_status, changed_at, changed_by)."
    )

    doc.add_paragraph().add_run("3.2. Модули db/database.py и db/import_json.py:").bold = True
    p_mod = doc.add_paragraph(
        "• db/database.py: функция get_connection() выполняет подключение и включает контроль внешних ключей PRAGMA foreign_keys = ON;\n"
        "• db/import_json.py: выполняет миграцию данных из JSON-файлов каталога data/ с помощью conn.executemany() и параметризации ?. "
        "Соблюдается порядок зависимостей: supervisors и students импортируются до topics. Реализован транзакционный контроль (commit/rollback)."
    )

    doc.add_paragraph().add_run("3.3. Разработка 17 SQL-запросов (sql/queries.sql):").bold = True
    sql_list = [
        "Выборка и сортировка: списки тем и руководителей (ORDER BY title/full_name), пагинация студентов (LIMIT 5 OFFSET 0);",
        "Фильтрация WHERE: выборка свободных тем (WHERE status='свободна'), поиск по ключевому слову (LIKE), дедлайн в интервале (BETWEEN), студенты без закрепленной темы (подзапрос NOT IN);",
        "Соединение таблиц JOIN: полная карточка темы с объединением трех таблиц (topics + supervisors + students), активные темы в работе, свободные темы преподавателя;",
        "Агрегатные функции: общий подсчет тем и условные суммы SUM(CASE WHEN status=... THEN 1 ELSE 0 END), показатели лимитов нагрузки (AVG, MIN, MAX);",
        "Группировка GROUP BY: распределение тем по руководителям (LEFT JOIN + GROUP BY для сохранения руководителей с 0 тем), темы по статусам, студенты по группам;",
        "Фильтрация групп HAVING: руководители с нагрузкой более 1 темы (HAVING COUNT(t.id) > 1), группы с числом студентов более 1."
    ]
    for s in sql_list:
        doc.add_paragraph(s, style='List Bullet')

    doc.add_paragraph().add_run("3.4. Интеграция с Django и веб-страницы:").bold = True
    p_dj = doc.add_paragraph(
        "В coursework/views.py и core/views.py настроена фабрика conn.row_factory = sqlite3.Row. "
        "Все пользовательские данные подставляются через параметры ? без f-строк. Реализованы сценарии:\n"
        "• /coursework/topics/ — каталог тем с параметризованным поиском LIKE;\n"
        "• /coursework/active/ — темы в работе (соединение topics, supervisors, students через JOIN);\n"
        "• /coursework/available/ — веб-форма поиска свободных тем с фильтрацией по дедлайну (подзапрос NOT IN);\n"
        "• /coursework/statistics/ — страница аналитики с агрегатами, GROUP BY и HAVING."
    )

    add_section_heading(doc, "4. Результат работы")
    doc.add_paragraph(
        "Создана реляционная база данных data/coursework.sqlite, файлы sql/schema.sql и sql/queries.sql (17 запросов), "
        "модули db/database.py и db/import_json.py. Реализованы 4 новые страницы веб-интерфейса, получающие "
        "данные из SQLite напрямую через стандартную библиотеку sqlite3."
    )

    add_section_heading(doc, "5. Контрольные проверки")
    doc.add_paragraph("• python manage.py check — 0 ошибок (System check identified no issues);\n"
                      "• pytest -v — 22 passed (включая автоматические тесты tests/test_pr7_db.py на схему, FK и все 17 SQL-запросов);\n"
                      "• flake8 . --extend-exclude=venv — 0 ошибок (соответствие PEP 8);\n"
                      "• python manage.py test — 21 passed (OK — модели и представления).")

    add_section_heading(doc, "6. Фиксация в Git")
    doc.add_paragraph("Коммит зафиксирован с меткой: git commit -m 'PR7: add SQLite database and SQL queries', создан тег git tag PR7.")

    add_section_heading(doc, "7. Вывод")
    doc.add_paragraph(
        "В ходе практической работы № 7 освоены проектирование реляционной базы данных SQLite, язык SQL, "
        "миграция данных из JSON и безопасное выполнение параметризованных запросов из Python через sqlite3. "
        "Данные успешно интегрированы в веб-интерфейс Django с использованием фабрики sqlite3.Row."
    )

    out_file = BASE_DIR / "Отчет_ПР7_Фролов_Кирилл_Иванович.docx"
    doc.save(str(out_file))
    print(f"Отчет ПР7 создан: {out_file}")


if __name__ == "__main__":
    generate_pr5_report()
    generate_pr6_report()
    generate_pr7_report()
