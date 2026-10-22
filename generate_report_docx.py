"""Скрипт генерации официального отчёта в формате DOCX по ПР5, ПР6, ПР7."""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent
DOCX_OUTPUT = BASE_DIR / "Отчет_Фролов_Кирилл_Иванович_ПР5_ПР6_ПР7.docx"

IMG1_PATH = Path(r"C:/Users/frolo/.gemini/antigravity/brain/d57d169d-cf3c-4c46-b631-cd3105e21cb4/.user_uploaded/media_1792668603168_644a1c2f.png")
IMG2_PATH = Path(r"C:/Users/frolo/.gemini/antigravity/brain/d57d169d-cf3c-4c46-b631-cd3105e21cb4/.user_uploaded/media_1792668612473_35f68279.png")


def set_cell_background(cell, hex_color):
    """Установить фоновый цвет ячейки таблицы."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Установить внутренние отступы в ячейке."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def create_document():
    doc = Document()

    # Настройки полей документа (2 см со всех сторон)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.79)
        section.right_margin = Inches(0.79)

    # Базовый стиль
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11.5)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # --- ТИТУЛЬНАЯ ЧАСТЬ / ШАПКА ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(
        "МИНОБРНАУКИ РОССИИ\n"
        "Федеральное государственное бюджетное образовательное учреждение высшего образования\n"
        "«МИРЭА — Российский технологический университет»\n"
        "Институт перспективных технологий и индустриального программирования\n"
        "Кафедра индустриального программирования\n\n"
    )
    title_run.font.size = Pt(10)
    title_run.font.bold = True

    h1_p = doc.add_paragraph()
    h1_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1_run = h1_p.add_run("ОТЧЕТ ПО ПРАКТИЧЕСКИМ РАБОТАМ № 5, 6, 7\n")
    h1_run.font.size = Pt(16)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(
        "по дисциплине: «Технологии разработки приложений на базе фреймворков»\n"
        "Сквозной индивидуальный проект: «Система учета тем курсовых работ»\n\n"
    )
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True

    # Таблица сведений об авторе
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Студент:", "Фролов Кирилл Иванович"),
        ("Учебная группа:", "ЭФБО 12-24"),
        ("Тема индивидуального проекта:", "Система учета тем курсовых работ"),
        ("Преподаватель:", "Ящун Татьяна Викторовна"),
        ("Семестр:", "5 семестр, 2026–2027 уч. год"),
    ]
    for idx, (label, val) in enumerate(info_data):
        row = info_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.5)
        cell_val.width = Inches(4.5)
        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_background(cell_val, "FAFAFA")
        set_cell_margins(cell_lbl, 80, 80, 120, 120)
        set_cell_margins(cell_val, 80, 80, 120, 120)

        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10.5)

        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- РАЗДЕЛ: МАТЕРИАЛЫ ЯНДЕКС ПРАКТИКУМА СО СКРИНШОТАМИ ---
    h_yp = doc.add_heading(level=1)
    r_yp = h_yp.add_run("1. ВЫПОЛНЕНИЕ МАТЕРИАЛОВ КУРСА «ЯНДЕКС ПРАКТИКУМ»")
    r_yp.font.size = Pt(13)
    r_yp.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    doc.add_paragraph(
        "В рамках освоения курса бэкенд-разработки на платформе Яндекс Практикум "
        "были в полном объёме выполнены теоретические и практические материалы "
        "двух ключевых разделов программы, предшествующих и сопровождающих практические работы 5, 6 и 7:"
    )

    # Таблица раздела 04
    doc.add_paragraph().add_run("Раздел 04. «Основы веб-разработки» (к ПР5 и ПР6):").bold = True
    t4 = doc.add_table(rows=6, cols=3)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers4 = ["№", "Тема курса Яндекс Практикум", "Статус освоения"]
    for c_idx, h_text in enumerate(headers4):
        cell = t4.rows[0].cells[c_idx]
        set_cell_background(cell, "0D3B66")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    rows4 = [
        ("1", "Создание проекта (структура Django-проекта, settings.py, manage.py)", "Выполнено (100%)"),
        ("2", "Пути и view-функции (протокол HTTP, планирование адресов, конвертеры путей)", "Выполнено (100%)"),
        ("3", "Вёрстка для бэкендера (HTML, структура документа, селекторы CSS, Bootstrap)", "Выполнено (100%)"),
        ("4", "HTML и шаблоны Django (теги extends/include/block, DTL, фильтры, статика)", "Выполнено (100%)"),
        ("5", "Финальный проект спринта: Блогикум (полный цикл веб-разработки)", "Выполнено (100%)"),
    ]
    for r_idx, row_vals in enumerate(rows4, 1):
        for c_idx, val in enumerate(row_vals):
            cell = t4.rows[r_idx].cells[c_idx]
            set_cell_background(cell, "F9FBFD" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x7A, 0x43)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Вставка скриншота 1
    if IMG1_PATH.exists():
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(str(IMG1_PATH), width=Inches(5.0))
        p_caption1 = doc.add_paragraph()
        p_caption1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c1 = p_caption1.add_run("Рисунок 1 — Подтверждение выполнения тем раздела «Основы веб-разработки» (Яндекс Практикум)")
        r_c1.font.italic = True
        r_c1.font.size = Pt(9.5)

    # Таблица раздела 05
    doc.add_paragraph().add_run("Раздел 05. «База данных и Django ORM» (к ПР7):").bold = True
    t5 = doc.add_table(rows=7, cols=3)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h_text in enumerate(["№", "Тема курса Яндекс Практикум", "Статус освоения"]):
        cell = t5.rows[0].cells[c_idx]
        set_cell_background(cell, "0D3B66")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    rows5 = [
        ("1", "Введение в базы данных (реляционная модель, SQL, выборка, JOIN, агрегация, группировка)", "Выполнено (100% к ПР7)"),
        ("2", "Отношения между таблицами", "Предусмотрено в ПР8"),
        ("3", "Django ORM. Модели", "Предусмотрено в ПР8-ПР9"),
        ("4", "Админ-зона Django", "Предусмотрено в ПР9"),
        ("5", "Django ORM. Получение информации из БД", "Предусмотрено в ПР9-ПР10"),
        ("6", "Финальный проект спринта: публикации для Блогикум", "Предусмотрено в ПР10"),
    ]
    for r_idx, row_vals in enumerate(rows5, 1):
        for c_idx, val in enumerate(row_vals):
            cell = t5.rows[r_idx].cells[c_idx]
            set_cell_background(cell, "F9FBFD" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if r_idx == 1 and c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x7A, 0x43)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Вставка скриншота 2
    if IMG2_PATH.exists():
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(str(IMG2_PATH), width=Inches(5.0))
        p_caption2 = doc.add_paragraph()
        p_caption2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c2 = p_caption2.add_run("Рисунок 2 — Подтверждение выполнения темы «Введение в базы данных» раздела «База данных и Django ORM»")
        r_c2.font.italic = True
        r_c2.font.size = Pt(9.5)

    # --- ПРАКТИЧЕСКАЯ РАБОТА № 5 ---
    h_pr5 = doc.add_heading(level=1)
    r_pr5 = h_pr5.add_run("2. ПРАКТИЧЕСКАЯ РАБОТА № 5: «DJANGO, HTTP И ОСНОВЫ ВЕБ-РАЗРАБОТКИ»")
    r_pr5.font.size = Pt(13)
    r_pr5.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    p_pr5_goal = doc.add_paragraph()
    p_pr5_goal.add_run("Цель работы: ").bold = True
    p_pr5_goal.add_run(
        "Изучить архитектуру веб-приложений на фреймворке Django, структуру конфигурационного пакета "
        "и приложений, принципы клиент-серверного взаимодействия по протоколу HTTP (запрос/ответ), "
        "настроить систему URL-маршрутизации и создать первые view-функции для предметной области "
        "«Система учета тем курсовых работ»."
    )

    doc.add_paragraph().add_run("Основные результаты выполнения ПР5:").bold = True
    pr5_points = [
        "Развернут конфигурационный пакет проекта coursework_project и два Django-приложения: core (общие и информационные страницы) и coursework (учет курсовых работ, тем, студентов и руководителей);",
        "Реализовано распределение URL-маршрутов на двух уровнях: корневой urls.py подключает маршруты приложений через include(), в приложениях используются конвертеры типов (<int:pk>);",
        "Настроен обработчик handler404 для возврата брендированной страницы в случае обращения к несуществующим ресурсам;",
        "Подключен CSS-фреймворк Bootstrap 5.3 через CDN, реализована адаптивная верстка страниц с навигационной панелью (navbar), карточками (card), бейджами статусов и адаптивными таблицами;",
        "Разработана интерактивная веб-страница http://127.0.0.1:8000/http-info/, наглядно демонстрирующая разбор элементов входящего HTTP-запроса (метод, путь, заголовки, query-параметры);",
        "Сохранена полная работоспособность консольной версии проекта ПР1–ПР3 (main.py, storage.py) и исходных JSON-файлов в каталоге data/."
    ]
    for pt in pr5_points:
        doc.add_paragraph(pt, style='List Bullet')

    # --- ПРАКТИЧЕСКАЯ РАБОТА № 6 ---
    h_pr6 = doc.add_heading(level=1)
    r_pr6 = h_pr6.add_run("3. ПРАКТИЧЕСКАЯ РАБОТА № 6: «ШАБЛОНЫ DJANGO И ПРОЕКТНАЯ РАЗРАБОТКА»")
    r_pr6.font.size = Pt(13)
    r_pr6.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    p_pr6_goal = doc.add_paragraph()
    p_pr6_goal.add_run("Цель работы: ").bold = True
    p_pr6_goal.add_run(
        "Освоить профессиональную систему шаблонизации Django Template Language (DTL). "
        "Исключить формирование HTML-разметки внутри Python-кода, разделив проект на базовый каркас "
        "и подключаемые модули. Настроить контекст, циклы, фильтры, именованные пространства имён "
        "и механизм подключения статических файлов."
    )

    doc.add_paragraph().add_run("Основные результаты выполнения ПР6:").bold = True
    pr6_points = [
        "Создан единый общепроектный базовый шаблон templates/base.html, содержащий разметку шапки, адаптивную навигацию, блок контента {% block content %} и подвал;",
        "Полностью устранены временные строковые генераторы HTML из Python-представлений (отказ от функции page() в пользу стандартного render(request, template, context));",
        "Реализованы специализированные шаблоны: topic_list.html (каталог тем), topic_detail.html (детальная карточка темы), student_list.html (список студентов), supervisor_list.html (руководители и индикаторы нагрузки);",
        "Внедрены конструкции DTL: перебор коллекций {% for %} с веткой обработки пустого списка {% empty %}, условные блоки {% if %}, фильтры данных (|length, |default, |date);",
        "Настроена именованная маршрутизация с пространствами имён (app_name = 'coursework', {% url 'coursework:topic_detail' topic.id %}), гарантирующая независимость шаблонов от изменений URL-адресов;",
        "Подключены собственные статические файлы проекта через {% load static %}: кастомная таблица стилей static/css/style.css, векторные иконки Bootstrap Icons."
    ]
    for pt in pr6_points:
        doc.add_paragraph(pt, style='List Bullet')

    # --- ПРАКТИЧЕСКАЯ РАБОТА № 7 ---
    h_pr7 = doc.add_heading(level=1)
    r_pr7 = h_pr7.add_run("4. ПРАКТИЧЕСКАЯ РАБОТА № 7: «ОСНОВЫ БАЗ ДАННЫХ И SQL»")
    r_pr7.font.size = Pt(13)
    r_pr7.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    p_pr7_goal = doc.add_paragraph()
    p_pr7_goal.add_run("Цель работы: ").bold = True
    p_pr7_goal.add_run(
        "Изучить реляционную модель данных и язык SQL в СУБД SQLite. Спроектировать структуру реляционных "
        "таблиц предметной области курсовых работ, определить первичные и внешние ключи, типы полей "
        "и ограничения целостности. Реализовать транзакционный скрипт импорта данных из JSON. "
        "Разработать каталог из 17 SQL-запросов и интегрировать их в Django-представления с использованием "
        "sqlite3.Row и параметризованных запросов."
    )

    doc.add_paragraph().add_run("4.1. Спроектированная структура таблиц (sql/schema.sql):").bold = True
    schema_points = [
        "students — студенты: id (INTEGER PRIMARY KEY), full_name, group_name, direction, email (UNIQUE), created_at (DEFAULT datetime('now'));",
        "supervisors — научные руководители: id (INTEGER PRIMARY KEY), full_name, degree, position, email (UNIQUE), max_topics (CHECK 1..20), created_at;",
        "topics — темы курсовых работ: id (INTEGER PRIMARY KEY), title, description, deadline, created_at, status (CHECK: 'свободна', 'закреплена', 'на проверке', 'защищена'), "
        "внешние ключи: supervisor_id REFERENCES supervisors(id) ON DELETE RESTRICT, student_id REFERENCES students(id) ON DELETE SET NULL;",
        "status_history — аудит изменения статусов: id (INTEGER PRIMARY KEY AUTOINCREMENT), topic_id REFERENCES topics(id) ON DELETE CASCADE, old_status, new_status, changed_at, changed_by."
    ]
    for pt in schema_points:
        doc.add_paragraph(pt, style='List Bullet')

    doc.add_paragraph().add_run("4.2. Модули подключения и импорта (db/database.py, db/import_json.py):").bold = True
    db_points = [
        "db/database.py: функция get_connection() выполняет обязательную команду PRAGMA foreign_keys = ON;, гарантируя целостность связей при всех операциях;",
        "db/import_json.py: читает исходные JSON-файлы каталога data/, использует безопасный транзакционный механизм (commit/rollback) и метод executemany() с параметризацией ?. "
        "Соблюдается строгая иерархия: независимые таблицы supervisors и students заполняются перед зависимой таблицей topics, исключая сбои ограничений внешних ключей."
    ]
    for pt in db_points:
        doc.add_paragraph(pt, style='List Bullet')

    doc.add_paragraph().add_run("4.3. Каталог 17 проверенных SQL-запросов (sql/queries.sql):").bold = True
    sql_types = [
        ("Выборка и сортировка (SELECT, ORDER BY, LIMIT, OFFSET)", "Запросы 1, 2, 3: алфавитные списки тем и преподавателей, пагинация студентов."),
        ("Фильтрация данных (WHERE)", "Запросы 4, 5, 6, 7: свободные темы (WHERE status='свободна'), параметризованный поиск по подстроке (LIKE), темы в диапазоне дат (BETWEEN), студенты без тем (подзапрос NOT IN)."),
        ("Соединение таблиц (INNER JOIN, LEFT JOIN)", "Запросы 8, 9, 10: полная сводка темы с объединением topics, supervisors и students, темы в работе, свободные темы конкретного преподавателя."),
        ("Агрегатные функции (COUNT, SUM, AVG, MIN, MAX)", "Запросы 11, 12: подсчет тем и условные суммы SUM(CASE WHEN status=... THEN 1 ELSE 0 END), показатели лимитов нагрузки преподавателей."),
        ("Группировка записей (GROUP BY)", "Запросы 13, 14, 15: количество тем по каждому руководителю (LEFT JOIN + GROUP BY для учета преподавателей без тем), распределение тем по статусам, студентов по группам."),
        ("Фильтрация сформированных групп (HAVING)", "Запросы 16, 17: преподаватели с нагрузкой более 1 темы (HAVING COUNT(t.id) > 1), группы с числом студентов более одного.")
    ]
    for title, desc in sql_types:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{title}: ").bold = True
        p.add_run(desc)

    doc.add_paragraph().add_run("4.4. Интеграция с веб-приложением Django:").bold = True
    django_sql_points = [
        "Во view-функциях coursework/views.py и core/views.py настроена фабрика строк conn.row_factory = sqlite3.Row, позволяющая обращаться к полям по именам столбцов в Python и шаблонах (row['title'], {{ topic.title }});",
        "Все значения передаются в SQL исключительно через заполнители ? (параметризованные запросы), что полностью исключает возможность SQL-инъекций;",
        "Реализованы новые пользовательские сценарии: /coursework/active/ (активные темы через JOIN), /coursework/available/ (поиск свободных тем по дате с подзапросом NOT IN), /coursework/statistics/ (дашборд со сводкой GROUP BY и HAVING)."
    ]
    for pt in django_sql_points:
        doc.add_paragraph(pt, style='List Bullet')

    # --- РАЗДЕЛ: ТЕСТИРОВАНИЕ И КОНТРОЛЬ КАЧЕСТВА ---
    h_test = doc.add_heading(level=1)
    r_test = h_test.add_run("5. РЕЗУЛЬТАТЫ КОНТРОЛЬНОГО ТЕСТИРОВАНИЯ И ПРОВЕРОК")
    r_test.font.size = Pt(13)
    r_test.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    doc.add_paragraph(
        "В соответствии с чек-листами практических работ были проведены все обязательные "
        "инструментальные проверки проекта:"
    )

    t_res = doc.add_table(rows=5, cols=3)
    t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h_text in enumerate(["Проверочный инструмент", "Команда терминала", "Результат выполнения"]):
        cell = t_res.rows[0].cells[c_idx]
        set_cell_background(cell, "0D3B66")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    test_rows = [
        ("Django System Check", "python manage.py check", "0 ошибок (System check identified no issues)"),
        ("Тестовый набор Pytest", "pytest -v", "22 passed (100% тестов пройдено за 0.03 с)"),
        ("Линтер стиля кода Flake8", "flake8 . --extend-exclude=venv", "0 замечаний (полное соответствие PEP 8)"),
        ("Набор тестов Django", "python manage.py test", "21 passed (OK — модели и представления)"),
    ]
    for r_idx, row_vals in enumerate(test_rows, 1):
        for c_idx, val in enumerate(row_vals):
            cell = t_res.rows[r_idx].cells[c_idx]
            set_cell_background(cell, "F9FBFD" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x7A, 0x43)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- РАЗДЕЛ: GIT ---
    h_git = doc.add_heading(level=1)
    r_git = h_git.add_run("6. ФИКСАЦИЯ ИЗМЕНЕНИЙ В СИСТЕМЕ КОНТРОЛЯ ВЕРСИЙ GIT")
    r_git.font.size = Pt(13)
    r_git.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    doc.add_paragraph(
        "Все этапы разработки проекта строго зафиксированы в локальном репозитории Git с соответствующими метками и коммитами:"
    )
    git_points = [
        "Контрольная точка ПР5: коммит с описанием переноса проекта на Django, создание метки (тега) git tag PR5;",
        "Контрольная точка ПР6: коммит PR6: add Django templates and static files — внедрение системы DTL и статических файлов;",
        "Контрольная точка ПР7: коммит PR7: add SQLite database and SQL queries (хэш e1489eb) и метка git tag PR7 — внедрение схемы SQLite, 17 запросов и новых view-функций."
    ]
    for pt in git_points:
        doc.add_paragraph(pt, style='List Bullet')

    # --- ЗАКЛЮЧЕНИЕ / ВЫВОД ---
    h_concl = doc.add_heading(level=1)
    r_concl = h_concl.add_run("7. ВЫВОД")
    r_concl.font.size = Pt(13)
    r_concl.font.color.rgb = RGBColor(0x0d, 0x3b, 0x66)

    doc.add_paragraph(
        "В результате последовательного выполнения практических работ № 5, 6 и 7 учебный сквозной "
        "индивидуальный проект «Система учета тем курсовых работ» успешно прошел путь от консольного прототипа "
        "до модульного веб-приложения на базе фреймворка Django. "
        "Были всесторонне освоены архитектурные принципы веб-разработки (MVT), протокол HTTP, маршрутизация, "
        "шаблонизация DTL с наследованием, стилизация на Bootstrap 5.3, а также реляционное моделирование данных, "
        "язык SQL и безопасная работа с СУБД SQLite через параметризованные запросы. "
        "Все тесты успешно проходят, код удовлетворяет стандарту PEP 8, проект полностью готов к защите."
    )

    doc.save(str(DOCX_OUTPUT))
    print(f"Отчет успешно создан: {DOCX_OUTPUT}")


if __name__ == "__main__":
    create_document()
